from __future__ import annotations

from dataclasses import dataclass, asdict
from io import BytesIO
from typing import Iterable
import re

import fitz  # PyMuPDF
from pptx import Presentation
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


@dataclass
class Chunk:
    text: str
    source: str
    page: int
    kind: str  # "lecture" or "exam"

    def to_dict(self):
        return asdict(self)


def _clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_uploaded_file(uploaded_file, kind: str) -> list[dict]:
    """Extract text while preserving file/page(or slide) metadata."""
    name = uploaded_file.name
    suffix = name.lower().rsplit(".", 1)[-1]
    raw = uploaded_file.getvalue()
    records: list[dict] = []

    if suffix == "pdf":
        doc = fitz.open(stream=raw, filetype="pdf")
        for i, page in enumerate(doc, start=1):
            text = _clean_text(page.get_text("text"))
            if text:
                records.append({"text": text, "source": name, "page": i, "kind": kind})

    elif suffix == "pptx":
        prs = Presentation(BytesIO(raw))
        for i, slide in enumerate(prs.slides, start=1):
            pieces = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    pieces.append(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        pieces.append(" | ".join(cell.text for cell in row.cells))
            text = _clean_text("\n".join(pieces))
            if text:
                records.append({"text": text, "source": name, "page": i, "kind": kind})

    elif suffix in {"txt", "md"}:
        text = raw.decode("utf-8", errors="ignore")
        text = _clean_text(text)
        if text:
            records.append({"text": text, "source": name, "page": 1, "kind": kind})

    else:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix}")

    return records


def chunk_records(records: Iterable[dict], chunk_size: int = 1500, overlap: int = 250) -> list[dict]:
    chunks: list[dict] = []
    for rec in records:
        text = rec["text"]
        if len(text) <= chunk_size:
            chunks.append(Chunk(**rec).to_dict())
            continue

        start = 0
        while start < len(text):
            end = min(len(text), start + chunk_size)
            piece = text[start:end]

            # 가능하면 문장/문단 경계에서 자르기
            if end < len(text):
                candidates = [
                    piece.rfind("\n\n"),
                    piece.rfind(". "),
                    piece.rfind("다. "),
                    piece.rfind("? "),
                ]
                cut = max(candidates)
                if cut > int(chunk_size * 0.55):
                    end = start + cut + 1
                    piece = text[start:end]

            piece = _clean_text(piece)
            if piece:
                chunks.append(
                    Chunk(
                        text=piece,
                        source=rec["source"],
                        page=rec["page"],
                        kind=rec["kind"],
                    ).to_dict()
                )

            if end >= len(text):
                break
            start = max(start + 1, end - overlap)

    return chunks


class LocalRetriever:
    """API 비용 없이 로컬에서 관련 강의/기출 chunk를 찾는 간단한 RAG retriever."""

    def __init__(self, chunks: list[dict]):
        self.chunks = chunks
        corpus = [c["text"] for c in chunks]
        # 한국어 + 수식/영문 약어에도 비교적 잘 작동하도록 문자 n-gram 사용
        self.vectorizer = TfidfVectorizer(
            analyzer="char_wb",
            ngram_range=(2, 5),
            max_features=35000,
            sublinear_tf=True,
        )
        self.matrix = self.vectorizer.fit_transform(corpus) if corpus else None

    def search(self, query: str, k: int = 6, kind: str | None = None) -> list[dict]:
        if not self.chunks or self.matrix is None:
            return []

        q = self.vectorizer.transform([query or "핵심 개념 시험 문제"])
        sims = cosine_similarity(q, self.matrix).ravel()
        ranked = sims.argsort()[::-1]

        result = []
        for idx in ranked:
            chunk = self.chunks[int(idx)]
            if kind and chunk["kind"] != kind:
                continue
            item = dict(chunk)
            item["score"] = float(sims[int(idx)])
            result.append(item)
            if len(result) >= k:
                break
        return result

    def representative(self, k: int = 8) -> list[dict]:
        if not self.chunks:
            return []
        if len(self.chunks) <= k:
            return list(self.chunks)
        # 문서 전체에 걸쳐 골고루 샘플
        step = (len(self.chunks) - 1) / max(k - 1, 1)
        idxs = sorted({round(i * step) for i in range(k)})
        return [self.chunks[i] for i in idxs]


def context_to_text(items: list[dict], max_chars: int = 12000) -> str:
    parts = []
    total = 0
    for i, item in enumerate(items, start=1):
        label = f"[자료 {i}] {item['source']} / {'슬라이드·페이지' if item['source'].lower().endswith('.pptx') else 'p.'}{item['page']} / {item['kind']}"
        block = f"{label}\n{item['text']}"
        if total + len(block) > max_chars:
            remain = max_chars - total
            if remain > 400:
                parts.append(block[:remain])
            break
        parts.append(block)
        total += len(block)
    return "\n\n".join(parts)


def reference_labels(items: list[dict]) -> list[str]:
    seen = set()
    labels = []
    for item in items:
        label = f"{item['source']} · {item['page']}p/slide"
        if label not in seen:
            seen.add(label)
            labels.append(label)
    return labels
